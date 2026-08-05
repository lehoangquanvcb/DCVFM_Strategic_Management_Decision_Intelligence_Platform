from __future__ import annotations

from io import BytesIO

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def board_pack_pdf(ex, market, regime, quality, advice, offtrack, actions) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("titlex", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18, textColor=colors.HexColor("#081426"), alignment=TA_CENTER, spaceAfter=12)
    h = ParagraphStyle("hx", parent=styles["Heading2"], fontName="Helvetica-Bold", textColor=colors.HexColor("#1F6FEB"), spaceBefore=8, spaceAfter=6)
    body = styles["BodyText"]
    story = [Paragraph("DCVFM Corporate Performance, Risk & Management Intelligence – Board Pack", title), Paragraph("V4.3 | Author: Le Hoang Quan", styles["Normal"]), Spacer(1, 8)]
    kpis = [["AUM", "Market regime", "Enterprise risk", "Off-track KPIs", "Data quality"], [f"{ex['aum']/1000:,.1f} tn VND", regime, f"{ex['risk_score']:.0f}/100", str(offtrack), f"{quality['score']:.0f}/100"]]
    t = Table(kpis, colWidths=[34*mm]*5)
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#081426")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("ALIGN",(0,0),(-1,-1),"CENTER"),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#CBD5E1")),("BOTTOMPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),7)]))
    story += [t, Spacer(1,10), Paragraph("Executive assessment", h), Paragraph(f"DCVFM is monitored at AUM {ex['aum']/1000:,.1f} trillion VND. Market regime is {regime}. {offtrack} full-year KPIs are off track and {actions['open']} management actions remain open. Data quality is {quality['score']:.0f}/100, with {quality['assumption_pct']:.0%} of required fields in assumption-led domains.", body), Paragraph("Top management recommendations", h)]
    rec_rows=[["Priority","Domain","Recommendation","Owner"]]
    for _,r in advice.head(6).iterrows(): rec_rows.append([str(r["Priority"]),str(r["Domain"]),Paragraph(str(r["Recommended_Action"]),body),str(r["Owner"])])
    rt=Table(rec_rows,colWidths=[20*mm,32*mm,95*mm,28*mm],repeatRows=1)
    rt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F6FEB")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#CBD5E1")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),8),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F1F5F9")])]))
    story += [rt, Paragraph("Caveat", h), Paragraph("Company, portfolio and operating rows marked Assumption must be verified before formal management, regulatory or investment use.", body)]
    doc.build(story)
    return buf.getvalue()


def board_pack_pptx(ex, regime, quality, advice, offtrack, actions) -> bytes:
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    navy=RGBColor(8,20,38); blue=RGBColor(31,111,235); white=RGBColor(255,255,255)
    def title_slide(title, subtitle=""):
        s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=navy
        box=s.shapes.add_textbox(Inches(.7),Inches(1.5),Inches(12),Inches(1.2)); p=box.text_frame.paragraphs[0]; p.text=title; p.font.size=Pt(30); p.font.bold=True; p.font.color.rgb=white
        sub=s.shapes.add_textbox(Inches(.72),Inches(3),Inches(11.5),Inches(.8)); q=sub.text_frame.paragraphs[0]; q.text=subtitle; q.font.size=Pt(15); q.font.color.rgb=RGBColor(143,163,191)
        return s
    title_slide("DCVFM Corporate Performance, Risk & Management Intelligence Platform", "Board Pack V4.3 | Author: Le Hoang Quan")
    s=title_slide("Executive control room", f"As of {ex['as_of']}")
    values=[("AUM",f"{ex['aum']/1000:,.1f} tn VND"),("Market regime",regime),("Enterprise risk",f"{ex['risk_score']:.0f}/100"),("Off-track KPIs",str(offtrack)),("Data quality",f"{quality['score']:.0f}/100")]
    for i,(label,value) in enumerate(values):
        x=.7+i*2.5; shp=s.shapes.add_shape(5,Inches(x),Inches(4.1),Inches(2.15),Inches(1.35)); shp.fill.solid(); shp.fill.fore_color.rgb=RGBColor(16,35,61); shp.line.color.rgb=blue
        tf=shp.text_frame; tf.text=label; tf.paragraphs[0].font.size=Pt(12); tf.paragraphs[0].font.color.rgb=RGBColor(143,163,191); p=tf.add_paragraph(); p.text=value; p.font.size=Pt(20); p.font.bold=True; p.font.color.rgb=white
    s=title_slide("Top management actions", "Observation → diagnosis → implication → accountable action")
    y=3.4
    for _,r in advice.head(5).iterrows():
        box=s.shapes.add_textbox(Inches(.8),Inches(y),Inches(11.8),Inches(.55)); p=box.text_frame.paragraphs[0]; p.text=f"{r['Priority']} | {r['Domain']}: {r['Recommended_Action']} — {r['Owner']}"; p.font.size=Pt(12); p.font.color.rgb=white; y+=.68
    s=title_slide("Operating execution posture", f"{actions['open']} open actions | {actions['avg_progress']:.0%} average progress")
    box=s.shapes.add_textbox(Inches(.8),Inches(3.3),Inches(11.7),Inches(2)); p=box.text_frame.paragraphs[0]; p.text="Immediate focus: close off-track business-plan KPIs, resolve operating and compliance exceptions, improve data verification, and assign accountable owners with deadlines."; p.font.size=Pt(18); p.font.color.rgb=white
    out=BytesIO(); prs.save(out); return out.getvalue()
